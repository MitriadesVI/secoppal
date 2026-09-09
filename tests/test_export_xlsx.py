from __future__ import annotations

from datetime import datetime
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from app.core.export_xlsx import build_results_xlsx, export_filename, xlsx_mime_type


def _load(payload: bytes):
    return load_workbook(BytesIO(payload))


def test_build_results_xlsx_has_expected_sheets_and_metadata():
    rows = [
        {
            "titulo": "Mantenimiento de vías",
            "entidad": "ALCALDÍA DE PRUEBA",
            "valor": "$1,500,000,000",
            "estado": "Abierto",
            "estado_op_label": "Recepción de ofertas",
            "fecha": "2026-08-28",
            "fecha_label": "Publicación",
            "url": "https://community.secop.gov.co/Public/Tendering/NoticeDetail/Index?noticeUID=CO1.NTC.1",
            "referencia_del_proceso": "LP-001-2026",
        }
    ]
    payload = build_results_xlsx(
        rows=rows,
        query="oportunidades de mantenimiento en Atlántico",
        resolved_params={
            "dataset": "procesos",
            "objeto": ["mantenimiento"],
            "departamento_resolved": "ATLÁNTICO",
            "valor_min": 500_000_000,
            "_internal": "no exportar",
        },
        dataset_id="p6dx-8zbt",
        total_count=17,
        generated_at=datetime(2026, 8, 28, 8, 15, tzinfo=ZoneInfo("America/Bogota")),
    )

    wb = _load(payload)
    assert wb.sheetnames == ["Resultados", "Búsqueda"]

    results = wb["Resultados"]
    assert results["A2"].value == "LP-001-2026"
    assert results["B2"].value == "Mantenimiento de vías"
    assert results["E2"].value == 1_500_000_000
    assert results["E2"].number_format == '$#,##0'
    assert results["F2"].value == "Recepción de ofertas"
    assert results["I2"].hyperlink.target.startswith("https://community.secop.gov.co/")

    search = wb["Búsqueda"]
    values = {search.cell(i, 1).value: search.cell(i, 2).value for i in range(2, search.max_row + 1)}
    assert values["Consulta original"] == "oportunidades de mantenimiento en Atlántico"
    assert values["Dataset SECOP"] == "p6dx-8zbt"
    assert values["Resultados exportados"] == 1
    assert values["Total reportado por la búsqueda"] == 17
    assert values["Filtro · departamento_resolved"] == "ATLÁNTICO"
    assert "Filtro · _internal" not in values


def test_formula_like_source_text_is_neutralized():
    payload = build_results_xlsx(
        rows=[
            {
                "titulo": "=HYPERLINK(\"https://evil.invalid\",\"click\")",
                "entidad": "+cmd",
                "valor": "$25,000,000",
                "estado": "Publicado",
                "fecha": "2026-08-28",
                "fecha_label": "Publicación",
                "url": "",
                "id_del_proceso": "CO1.REQ.123",
            }
        ],
        query="=1+1",
        resolved_params={"objeto": ["=SUM(A1:A2)"]},
        dataset_id="p6dx-8zbt",
    )
    wb = _load(payload)
    results = wb["Resultados"]
    search = wb["Búsqueda"]
    assert results["B2"].value.startswith("'=")
    assert results["C2"].value.startswith("'+")
    assert search["B2"].value.startswith("'=")


def test_empty_rows_still_build_valid_workbook():
    from app.core.contract_universe import stamp
    payload = build_results_xlsx(
        rows=[],
        query="contratos de seguros",
        resolved_params=stamp({"dataset": "contratos"}),
        dataset_id="jbjy-vk9h",
        total_count=0,
    )
    wb = _load(payload)
    assert wb["Resultados"].max_row == 1
    assert wb["Búsqueda"]["B5"].value == 0


def test_filename_and_mime_type_are_stable():
    instant = datetime(2026, 8, 28, 7, 45, tzinfo=ZoneInfo("America/Bogota"))
    assert export_filename(instant) == "secoppal_resultados_20260828_0745.xlsx"
    assert xlsx_mime_type() == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

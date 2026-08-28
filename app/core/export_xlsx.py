"""XLSX export for already-rendered SECOPPAL search results.

V1 is intentionally presentation-only: it never queries SECOP, rebuilds SoQL,
or expands the visible result universe. Callers pass the rows already shown to
the user plus the interpreted search metadata.
"""
from __future__ import annotations

import json
import re
from datetime import datetime
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_COLOMBIA_TZ = ZoneInfo("America/Bogota")
_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_ILLEGAL_XML_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_FORMULA_PREFIXES = ("=", "+", "-", "@")

# User-facing filters only. Avoid exporting resolver/debug payloads that happen
# to live inside resolved_params.
_FILTER_KEYS = (
    "dataset",
    "intent_type",
    "objeto",
    "objeto_or",
    "objeto_phrases",
    "entidad_resolved",
    "departamento_resolved",
    "ciudad",
    "contratista",
    "actor_identity",
    "modalidad",
    "estado_family",
    "estado",
    "fecha_desde",
    "fecha_hasta",
    "valor_min",
    "valor_max",
    "ordering_signal",
)

_RESULT_COLUMNS = (
    ("Referencia", "reference"),
    ("Título / objeto", "titulo"),
    ("Entidad", "entidad"),
    ("Contratista", "contratista"),
    ("Valor (COP)", "valor"),
    ("Estado", "estado_visible"),
    ("Fecha", "fecha"),
    ("Tipo de fecha", "fecha_label"),
    ("Enlace SECOP", "url"),
)


def xlsx_mime_type() -> str:
    return _XLSX_MIME


def export_filename(generated_at: datetime | None = None) -> str:
    instant = _colombia_time(generated_at)
    return f"secoppal_resultados_{instant:%Y%m%d_%H%M}.xlsx"


def build_results_xlsx(
    *,
    rows: list[dict],
    query: str,
    resolved_params: dict | None,
    dataset_id: str,
    total_count: int = 0,
    generated_at: datetime | None = None,
) -> bytes:
    """Build a two-sheet workbook from the rows already visible to the user."""
    instant = _colombia_time(generated_at)
    workbook = Workbook()
    results_ws = workbook.active
    results_ws.title = "Resultados"
    search_ws = workbook.create_sheet("Búsqueda")

    _write_results_sheet(results_ws, rows)
    _write_search_sheet(
        search_ws,
        query=query,
        resolved_params=resolved_params or {},
        dataset_id=dataset_id,
        exported_count=len(rows),
        total_count=total_count,
        instant=instant,
    )

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _colombia_time(value: datetime | None) -> datetime:
    if value is None:
        return datetime.now(_COLOMBIA_TZ)
    if value.tzinfo is None:
        return value.replace(tzinfo=_COLOMBIA_TZ)
    return value.astimezone(_COLOMBIA_TZ)


def _safe_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, (dict, list, tuple, set)):
        if isinstance(value, set):
            value = sorted(value)
        text = json.dumps(value, ensure_ascii=False, sort_keys=True, default=str)
    else:
        text = str(value)
    text = _ILLEGAL_XML_RE.sub("", text)
    if text.startswith(_FORMULA_PREFIXES):
        return "'" + text
    return text


def _numeric_cop(value: object) -> int | float | str:
    """Turn Formatter's '$1,234' display value back into a sortable Excel number."""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return value
    text = str(value or "").strip()
    if not text or text.lower() == "sin valor":
        return ""
    cleaned = text.replace("$", "").replace(",", "").strip()
    try:
        number = float(cleaned)
    except ValueError:
        return _safe_text(text)
    return int(number) if number.is_integer() else number


def _reference(row: dict) -> str:
    return _safe_text(
        row.get("referencia_del_proceso")
        or row.get("referencia_del_contrato")
        or row.get("id_del_proceso")
        or row.get("id_contrato")
        or ""
    )


def _visible_state(row: dict) -> str:
    return _safe_text(row.get("estado_op_label") or row.get("estado") or "")


def _result_value(row: dict, key: str) -> object:
    if key == "reference":
        return _reference(row)
    if key == "estado_visible":
        return _visible_state(row)
    if key == "valor":
        return _numeric_cop(row.get("valor"))
    return _safe_text(row.get(key, ""))


def _write_results_sheet(ws, rows: list[dict]) -> None:
    headers = [label for label, _ in _RESULT_COLUMNS]
    ws.append(headers)

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="top", wrap_text=True)

    for row in rows:
        ws.append([_result_value(row, key) for _, key in _RESULT_COLUMNS])

    value_col = next(i for i, (_, key) in enumerate(_RESULT_COLUMNS, start=1) if key == "valor")
    url_col = next(i for i, (_, key) in enumerate(_RESULT_COLUMNS, start=1) if key == "url")
    for row_idx in range(2, ws.max_row + 1):
        ws.cell(row=row_idx, column=value_col).number_format = '$#,##0'
        url_cell = ws.cell(row=row_idx, column=url_col)
        url = url_cell.value
        if isinstance(url, str) and url.startswith(("http://", "https://")):
            url_cell.hyperlink = url
            url_cell.style = "Hyperlink"

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    _set_widths(ws, (20, 48, 34, 30, 18, 24, 16, 20, 45))
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def _write_search_sheet(
    ws,
    *,
    query: str,
    resolved_params: dict,
    dataset_id: str,
    exported_count: int,
    total_count: int,
    instant: datetime,
) -> None:
    ws.append(["Campo", "Valor"])
    metadata = [
        ("Consulta original", _safe_text(query)),
        ("Dataset SECOP", _safe_text(dataset_id)),
        ("Generado en", instant.isoformat(timespec="seconds")),
        ("Resultados exportados", exported_count),
        ("Total reportado por la búsqueda", int(total_count or 0)),
    ]
    for label, value in metadata:
        ws.append([label, value])

    for key in _FILTER_KEYS:
        value = resolved_params.get(key)
        if value in (None, "", [], {}, ()):
            continue
        ws.append([f"Filtro · {key}", _safe_text(value)])

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    ws.freeze_panes = "A2"
    _set_widths(ws, (34, 90))
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def _set_widths(ws, widths: tuple[int, ...]) -> None:
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
